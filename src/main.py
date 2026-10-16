"""Главный модуль, связывающий функциональности проекта."""

import csv
import json
from typing import Any

from src.operations import process_bank_operations
from src.search import process_bank_search

VALID_STATUSES = ("EXECUTED", "CANCELED", "PENDING")


def _load_json(path: str) -> list[dict[str, Any]]:
    """Загружает операции из JSON-файла."""
    with open(path, encoding="utf-8") as file:
        return json.load(file)


def _load_csv(path: str) -> list[dict[str, Any]]:
    """Загружает операции из CSV-файла."""
    with open(path, encoding="utf-8") as file:
        return list(csv.DictReader(file))


def _load_xlsx(path: str) -> list[dict[str, Any]]:
    """Загружает операции из XLSX-файла."""
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    headers = rows[0]
    return [dict(zip(headers, row)) for row in rows[1:]]


def _ask_yes_no(prompt: str) -> bool:
    """Задаёт пользователю вопрос да/нет и возвращает True для 'да'."""
    while True:
        answer = input(prompt).strip().lower()
        if answer in ("да", "нет"):
            return answer == "да"
        print("Пожалуйста, введите 'да' или 'нет'.")


def _filter_by_status(
    operations: list[dict[str, Any]], status: str
) -> list[dict[str, Any]]:
    """Фильтрует операции по статусу (без учёта регистра)."""
    return [
        operation
        for operation in operations
        if str(operation.get("state", "")).upper() == status
    ]


def _sort_by_date(
    operations: list[dict[str, Any]], ascending: bool
) -> list[dict[str, Any]]:
    """Сортирует операции по дате."""
    return sorted(
        operations,
        key=lambda operation: str(operation.get("date", "")),
        reverse=not ascending,
    )


def _filter_rubles(operations: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Оставляет только операции в рублях."""
    return [
        operation
        for operation in operations
        if str(operation.get("currency", "")).upper() == "RUB"
        or "руб" in str(operation.get("amount", ""))
    ]


def _print_operations(operations: list[dict[str, Any]]) -> None:
    """Печатает итоговый список операций."""
    if not operations:
        print(
            "Не найдено ни одной транзакции, подходящей под ваши "
            "условия фильтрации"
        )
        return

    print(f"Всего банковских операций в выборке: {len(operations)}")
    for operation in operations:
        print()
        print(f"{operation.get('date', '')} {operation.get('description', '')}")
        print(operation.get("from", ""))
        print(f"Сумма: {operation.get('amount', '')}")


def main() -> None:
    """Запускает CLI-программу работы с банковскими транзакциями."""
    print(
        "Привет! Добро пожаловать в программу работы "
        "с банковскими транзакциями."
    )
    print("Выберите необходимый пункт меню:")
    print("1. Получить информацию о транзакциях из JSON-файла")
    print("2. Получить информацию о транзакциях из CSV-файла")
    print("3. Получить информацию о транзакциях из XLSX-файла")

    choice = input().strip()
    if choice == "1":
        print("Для обработки выбран JSON-файл.")
        operations = _load_json("data/operations.json")
    elif choice == "2":
        print("Для обработки выбран CSV-файл.")
        operations = _load_csv("data/transactions.csv")
    else:
        print("Для обработки выбран XLSX-файл.")
        operations = _load_xlsx("data/transactions.xlsx")

    while True:
        print(
            "Введите статус, по которому необходимо выполнить фильтрацию. "
            f"Доступные для фильтровки статусы: {', '.join(VALID_STATUSES)}"
        )
        status = input().strip().upper()
        if status in VALID_STATUSES:
            print(f'Операции отфильтрованы по статусу "{status}"')
            break
        print(f'Статус операции "{status}" недоступен.')

    operations = _filter_by_status(operations, status)

    if _ask_yes_no("Отсортировать операции по дате? Да/Нет"):
        order = input(
            "Отсортировать по возрастанию или по убыванию? "
        ).strip().lower()
        operations = _sort_by_date(operations, ascending="возраст" in order)

    if _ask_yes_no("Выводить только рублевые транзакции? Да/Нет"):
        operations = _filter_rubles(operations)

    if _ask_yes_no(
        "Отфильтровать список транзакций по определенному слову "
        "в описании? Да/Нет"
    ):
        word = input("Введите слово для поиска: ").strip()
        operations = process_bank_search(operations, word)

    print("Распечатываю итоговый список транзакций...")
    _print_operations(operations)


if __name__ == "__main__":
    main()